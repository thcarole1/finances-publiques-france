"""Listes nationales des impôts (DG TAXUD), dernier instantané, au format long.

Une ligne par pays × impôt national × année, avec son code SEC 2010 (sto), sa base économique
(fonction) et son montant en millions de monnaie nationale. Les lignes de mémo (fonction _Z,
Z ou vide) sont conservées : le filtrage se fait dans les modèles intermédiaires.
"""

from pathlib import Path

import openpyxl
import pandas as pd

RACINE = Path("data/raw/taxud/ntl")
FONCTIONS_LIGNES_PART = {"L", "S"}  # lignes de répartition SPLIT2 (parts, pas des montants)


def dernier_instantane() -> tuple[Path, str]:
    dossiers = sorted(p for p in RACINE.iterdir() if p.is_dir())
    if not dossiers:
        raise FileNotFoundError("Aucun instantané DG TAXUD dans data/raw/taxud/ntl")
    return dossiers[-1] / "ntl.xlsx", dossiers[-1].name


def nombre(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def model(dbt, session):
    dbt.config(materialized="table")
    fichier, date_collecte = dernier_instantane()
    classeur = openpyxl.load_workbook(fichier, read_only=True, data_only=True)
    lignes = []
    for feuille in classeur.sheetnames:
        if len(feuille) != 2:  # feuilles pays : codes à deux lettres
            continue
        rangees = list(classeur[feuille].iter_rows(values_only=True))
        i0 = next(i for i, r in enumerate(rangees) if r[0] and str(r[0]).startswith("STO"))
        entete = rangees[i0]
        annees = {j: int(c) for j, c in enumerate(entete) if isinstance(c, int) and 1990 < c < 2100}
        for r in rangees[i0 + 1 :]:
            fonction = r[9] if len(r) > 9 else None
            if not r[0] or fonction in FONCTIONS_LIGNES_PART:
                continue
            for j, annee in annees.items():
                lignes.append(
                    {
                        "geo": feuille,
                        "annee": annee,
                        "sto": str(r[0]).strip(),
                        "details": str(r[1]).strip() if r[1] is not None else None,
                        "libelle_fr": r[2],
                        "libelle_en": r[5],
                        "fonction": str(fonction).strip() if fonction is not None else None,
                        "valeur": nombre(r[j]),
                        "date_collecte": date_collecte,
                    }
                )
    df = pd.DataFrame(lignes)
    df["annee"] = df["annee"].astype("int32")
    df["date_collecte"] = pd.to_datetime(df["date_collecte"]).dt.date
    return df
