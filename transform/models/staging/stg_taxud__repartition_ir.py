"""Répartitions déclarées par les pays (DG TAXUD), dernier instantané, au format long.

- SPLIT1 : impôt sur le revenu des personnes physiques réparti entre travail (LEES),
  indépendants (KISE), capital (KIH) et revenus de remplacement (LNON) — feuille « PIT split » ;
- SPLIT2 : prélèvements répartis entre travail (LEES) et indépendants (KISE) — lignes L et S
  en bas de chaque feuille pays.
"""

from pathlib import Path

import openpyxl
import pandas as pd

RACINE = Path("data/raw/taxud/ntl")
SPLIT1 = {"labour": "LEES", "self-employed": "KISE", "capital": "KIH", "transfers": "LNON"}
SPLIT2 = {"L": "LEES", "S": "KISE"}


def model(dbt, session):
    dbt.config(materialized="table")
    dossier = sorted(p for p in RACINE.iterdir() if p.is_dir())[-1]
    classeur = openpyxl.load_workbook(dossier / "ntl.xlsx", read_only=True, data_only=True)
    lignes = []

    feuille_pit = next(n for n in classeur.sheetnames if "PIT" in n.upper())
    rangees = list(classeur[feuille_pit].iter_rows(values_only=True))
    i0 = next(i for i, r in enumerate(rangees) if r[0] == "GEO")
    annees = {j: int(c) for j, c in enumerate(rangees[i0]) if isinstance(c, int)}
    for r in rangees[i0 + 1 :]:
        if r[0] and r[1] in SPLIT1:
            for j, annee in annees.items():
                if isinstance(r[j], (int, float)):
                    lignes.append(("SPLIT1", r[0], annee, SPLIT1[r[1]], float(r[j])))

    for feuille in classeur.sheetnames:
        if len(feuille) != 2:
            continue
        rangees = list(classeur[feuille].iter_rows(values_only=True))
        i0 = next(i for i, r in enumerate(rangees) if r[0] and str(r[0]).startswith("STO"))
        annees = {
            j: int(c) for j, c in enumerate(rangees[i0]) if isinstance(c, int) and 1990 < c < 2100
        }
        for r in rangees[i0 + 1 :]:
            if len(r) > 10 and r[9] in SPLIT2 and str(r[10] or "").startswith(feuille):
                for j, annee in annees.items():
                    if isinstance(r[j], (int, float)):
                        lignes.append(("SPLIT2", feuille, annee, SPLIT2[r[9]], float(r[j])))

    df = pd.DataFrame(lignes, columns=["repartition", "geo", "annee", "fonction", "part"])
    df["annee"] = df["annee"].astype("int32")
    return df
